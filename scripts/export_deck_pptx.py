from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

CHROME = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
DECK = Path(__file__).resolve().parent.parent / "docs" / "presentation" / "deck.html"
DEMO_STEPS = 4  # slide with id="demo" is captured once per step (0..3)
WINDOW = "1920,1080"

EMU_PER_IN = 914400
SLIDE_W_IN = 13.333  # 16:9
SLIDE_H_IN = 7.5


def capture_targets(deck_html: str) -> list[tuple[int, int | None]]:
    """Build the capture list from the deck markup.

    Args:
        deck_html: Raw contents of deck.html.

    Returns:
        List of (slide_number, demo_step) pairs in presentation order;
        demo_step is None for ordinary slides.

    Raises:
        ValueError: If no slides are found in the markup.
    """
    sections = re.findall(r'<section class="slide"[^>]*>', deck_html)
    if not sections:
        raise ValueError("no <section class=\"slide\"> blocks found in deck.html")
    targets: list[tuple[int, int | None]] = []
    for i, tag in enumerate(sections, start=1):
        if 'id="demo"' in tag:
            targets.extend((i, step) for step in range(DEMO_STEPS))
        else:
            targets.append((i, None))
    return targets


def screenshot(slide: int, step: int | None, out_png: Path) -> None:
    """Render one deck slide to a PNG with headless Chrome.

    Args:
        slide: 1-based slide number (the deck's #N hash).
        step: Demo step to preset via ?step=N, or None for ordinary slides.
        out_png: Destination PNG path.

    Raises:
        RuntimeError: If Chrome exits non-zero or writes no file.
    """
    query = f"?step={step}" if step is not None else ""
    url = f"file://{DECK}{query}#{slide}"
    cmd = [
        str(CHROME), "--headless=new", "--disable-gpu", "--hide-scrollbars",
        f"--window-size={WINDOW}", "--force-device-scale-factor=2",
        "--virtual-time-budget=3000", f"--screenshot={out_png}", url,
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0 or not out_png.exists():
        raise RuntimeError(f"chrome screenshot failed for {url}: {result.stderr[-400:]}")


def build_pptx(pngs: list[Path], out_pptx: Path) -> None:
    """Assemble full-bleed image slides into a 16:9 .pptx.

    Args:
        pngs: Slide images in presentation order.
        out_pptx: Destination .pptx path.
    """
    from pptx import Presentation  # deferred: run via `uv run --with python-pptx`

    prs = Presentation()
    prs.slide_width = int(SLIDE_W_IN * EMU_PER_IN)
    prs.slide_height = int(SLIDE_H_IN * EMU_PER_IN)
    blank = prs.slide_layouts[6]
    for png in pngs:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(str(out_pptx))


def main() -> int:
    """Export deck.html to a .pptx of rendered slides.

    Returns:
        Process exit code.
    """
    parser = argparse.ArgumentParser(description="Render deck.html to a .pptx")
    parser.add_argument("-o", "--out", type=Path, default=Path("deck.pptx"),
                        help="output .pptx path (default: ./deck.pptx)")
    args = parser.parse_args()

    if not CHROME.exists():
        print(f"Chrome not found at {CHROME}", file=sys.stderr)
        return 1
    targets = capture_targets(DECK.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory() as tmp:
        pngs: list[Path] = []
        for slide, step in targets:
            suffix = f"-s{step}" if step is not None else ""
            png = Path(tmp) / f"slide{slide:02d}{suffix}.png"
            screenshot(slide, step, png)
            pngs.append(png)
            print(f"captured #{slide}{suffix} -> {png.name}")
        build_pptx(pngs, args.out)
    print(f"wrote {args.out} ({len(targets)} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
